from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from copy import deepcopy
from lxml import etree as E
import json,io
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parents[1]
SRC=R/'07_GestureCheckDraft_reviewed_tracked_CLAUDE_20_10_2026.docx'
OUT=R/'07_GestureCheck_Final_20_10_2026.docx'
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
W=NS['w']; Q=lambda s:'{'+W+'}'+s
with SRC.open('rb') as f,ZipFile(f) as z: parts={i.filename:z.read(i.filename) for i in z.infolist()}
root=E.fromstring(parts['word/document.xml']);ps=root.xpath('//w:body//w:p',namespaces=NS)
# Accept the reviewed draft's text and formatting revisions throughout every story.
for name,data in list(parts.items()):
 if not name.endswith('.xml') or not name.startswith('word/'):continue
 try:tree=root if name=='word/document.xml' else E.fromstring(data)
 except E.XMLSyntaxError:continue
 for el in reversed(tree.xpath('//w:del | //w:moveFrom | //w:pPrChange | //w:rPrChange | //w:tblPrChange | //w:tcPrChange | //w:trPrChange | //w:sectPrChange | //w:commentRangeStart | //w:commentRangeEnd | //w:commentReference | //w:trackRevisions',namespaces=NS)):
  el.getparent().remove(el)
 for el in reversed(tree.xpath('//w:ins | //w:moveTo',namespaces=NS)):
  parent=el.getparent();idx=parent.index(el);children=list(el);parent.remove(el)
  for j,ch in enumerate(children):parent.insert(idx+j,ch)
 parts[name]=E.tostring(tree,xml_declaration=True,encoding='UTF-8',standalone=True)

def text(i):return ''.join(ps[i].xpath('.//w:t/text()',namespaces=NS))
def replace(i,s,bold_lead=False):
 p=ps[i];fmt=p.find('w:pPr',NS);old=p.find('.//w:rPr',NS)
 for ch in list(p):
  if ch is not fmt:p.remove(ch)
 chunks=[s]
 if bold_lead and '). ' in s:pos=s.index('). ')+2;chunks=[s[:pos],s[pos:]]
 for j,chunk in enumerate(chunks):
  run=E.SubElement(p,Q('r'))
  if old is not None:
   rp=deepcopy(old)
   for b in rp.findall('w:b',NS)+rp.findall('w:bCs',NS):rp.remove(b)
   run.append(rp)
  else:rp=E.SubElement(run,Q('rPr'))
  if bold_lead and j==0:E.SubElement(rp,Q('b'))
  t=E.SubElement(run,Q('t'));t.set('{http://www.w3.org/XML/1998/namespace}space','preserve');t.text=chunk
 return p
def add_after(i,s):
 p=deepcopy(ps[i]);
 for ch in list(p):
  if ch.tag!=Q('pPr'):p.remove(ch)
 rr=E.SubElement(p,Q('r'));tt=E.SubElement(rr,Q('t'));tt.text=s
 ps[i].addnext(p);return p
def delete(i):ps[i].getparent().remove(ps[i])

replace(136,text(136).replace('associated with Macedonia et al. (2011) and Mayer et al. (2015).','whose collection description requests citation of Macedonia, Müller, and Friederici (2011) and Mayer et al. (2015).').replace('The 44-clip collection comprises original video stimuli from Macedonia, Müller, and Friederici (2011).','The separately supplied 44-clip collection also comprises original stimuli attributed to Macedonia, Müller, and Friederici (2011); the shared citation does not imply that the two collections contain identical files.'))
replace(138,'Each clip shows a female performer against a neutral background from a fixed frontal camera position. An ffprobe audit of the 134 MP4 presentation files found that the 90-clip gesture collection had a mean duration of 5.31 s (range 4.16–5.96), resolution 960 × 538 pixels, and frame rate 25 fps. The 44-clip object collection had a mean duration of 4.99 s (range 4.48–5.20); 43 files were 500 × 720 pixels and one was 960 × 538 pixels, all at 25 fps. None contained an audio stream. These silent MP4 files supplied the human presentation and model-input pipeline; GPT conditions received extracted JPEG frames. These measurements characterize the distributed presentation files rather than the original camera recordings. Performer identity across the two source collections was not independently verified.')
replace(261,'The video-centered omnibus analyses indicated cohort differences in every dimension (all p < .001). Table 3 reports all 21 video-adjusted, participant-clustered cohort contrasts, with pointwise 95% confidence intervals and Holm-adjusted p values across the full family of 21 tests. Seven contrasts had unadjusted p < .05, but only two survived Holm correction: the English-language cohort rated facial-expression emotional salience lower than the German-language cohort (b = −0.41, 95% CI [−0.60, −0.21], adjusted p = .001), and the Japanese-language cohort rated cultural familiarity lower (b = −0.31, 95% CI [−0.48, −0.14], adjusted p = .008). The remaining pairwise contrasts are descriptive exploratory estimates rather than multiplicity-adjusted evidence of a difference.')
replace(295,'Note. All 21 contrasts from seven ordinary least squares models are shown. German is the reference cohort; each model includes fixed effects for all 148 pairings and standard errors clustered by participant. Confidence intervals are pointwise 95% intervals using t(229); p values use the same reference distribution. Holm adjustment controls the family-wise error rate across all 21 contrasts.')
replace(296,'Response-style diagnostics were computed separately for each dimension and cohort as the proportions of ratings at either endpoint (1 or 5) and at the midpoint (3); Appendix A reports all values. For enactment potential, the Japanese cohort used endpoints less often (34.3% versus 39.5–48.0% in the other cohorts) and the midpoint more often (25.4% versus 16.1–17.2%). Cultural-familiarity endpoint use was 42.4% in the Japanese cohort versus 36.1–49.2% elsewhere, with midpoint use 18.8% versus 16.7–22.8%. These descriptive distributions do not isolate response style from substantive interpretation; the item-adjusted cohort effects should not be interpreted as pure cultural effects.')
replace(298,'These results indicate that at least some cohort differences persist when identical stimuli are compared, with the strongest multiplicity-adjusted evidence concerning English-language facial-expression ratings and Japanese-language cultural-familiarity ratings. Differences may reflect interpretation, scale use, translation, or a combination of these factors. Pooled human ratings are therefore population-dependent reference values. Cohort-specific validation is especially relevant for cultural familiarity, while the other exploratory contrasts require replication.')
replace(300,text(300)+' Appendix B records the exact saved model identifiers and access routes, including the auxiliary frame condition. Exact evaluation dates were not preserved in the retained records for the earlier Gemini, GPT, and Qwen runs; Glimpse was evaluated on 19 October 2026. The native 2-fps Gemini runs and the auxiliary Flash frame run also lack logged exact model identifiers. Model family is confounded with input route and preprocessing (native video, extracted frames, and provider-default sampling), so the benchmark does not isolate an architecture effect.')
replace(303,text(303).replace('the video itself','the video or sampled frames'))
replace(305,text(305).replace(', and a self-reported confidence check','').replace(', a short justification per dimension',', and a short justification per dimension'))
parsing='Parsing and field normalization. Saved model scores were checked against the seven canonical dimension names and the integer range 1–5 before evaluation. Alternate Qwen Max fields named motional_salience and facial_emotional_salience were mapped to the corresponding canonical dimensions. The normalization log records 127 field substitutions across 125 of 149 pairings (83.9%) for Qwen Max and none for the other configurations. This corrected incomplete field extraction rather than changing the underlying numeric responses: for motional salience, the correlation changed from −.20 on 22 standard pairings to +.27 on all 134. All configurations have complete final vectors for all 149 model pairings (0/149 missing final vectors per configuration). Appendix B distinguishes this final completeness check from historical request-level parsing failures: complete failure and retry histories were not retained for all configurations, so an initial parse-failure rate cannot be reconstructed.'
add_after(305,parsing);delete(495)
replace(313,text(313).replace('Section 6.3.2','Section 6.1'))
replace(613,'RQ1/H1 (correspondence with human ratings). H1 predicted positive correlations between model and human ratings across dimensions. The observed correlations support positive overall correspondence, with an important qualification: most configurations showed positive bias, with Glimpse the exception, so agreement in ranking did not imply agreement in absolute level.',True)
replace(655,'Cultural familiarity was among the largest sources of model–human error. Its human reference is population-dependent: after video adjustment and Holm correction across 21 contrasts, the Japanese-language cohort rated cultural familiarity lower than the German-language cohort, and the English-language cohort rated facial-expression emotional salience lower. Other cohort contrasts did not survive this correction. The response-distribution diagnostics in Appendix A suggest that scale use may also contribute, particularly for enactment potential, but do not separate scale use from interpretation. A further task mismatch is that humans were asked about familiarity “in your country,” whereas models were asked about “your cultural context,” without a defined reference population. This could contribute to model–human differences and was not experimentally isolated. A targeted rerun specifying the same reference population for models and humans—for example, typical adults in Germany—would test this explanation; no such rerun is claimed here. Cohort-specific validation and an explicit intended reference population are needed before these ratings are transferred to a different instructional or cultural setting.')
delete(662)
replace(663,text(663).replace('Second,','Several limitations qualify these findings. First,'))
replace(664,'Second, the human reference is population-dependent. The strongest adjusted cohort evidence concerns cultural familiarity and facial-expression salience; the remaining pairwise differences are exploratory. Model error should be read as error against the sampled cohort composition, and the response-style diagnostics do not establish why cohort distributions differ. Equal coverage across items and cohorts should not be assumed.')
replace(665,text(665).replace('Fourth,','Third,').replace('Relatedly, all gesture videos featured a single performer recorded once under standardized conditions, so whether model or human ratings would generalize to a different performer, a different recording style, or repeated takes of the same gesture is not addressed by this design.','The two source collections differ in framing and file geometry, and performer identity across collections was not independently verified. The design does not test generalization across performers, recording styles, or repeated takes.'))
replace(666,text(666).replace('Fifth,','Fourth,'))
replace(669,text(669).replace('A reliability study, repeating inference on the same items across multiple runs, prompt paraphrases, and, where available, temperature or sampling settings, would establish whether VLM ratings are stable enough to trust a single pass, which the present design cannot determine. ',''))
replace(673,text(673)+' The linked release predates the additional Holm correction and response-distribution diagnostics reported here. Those results are included in Table 3 and Appendix A of this manuscript. The controlled participant workbook is not redistributed. Any future trial-level release requires confirmation that consent and de-identification permit it; no public trial-level archive or archival DOI is claimed.')

parts['word/document.xml']=E.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
# Strip all comment parts and relationship/content-type entries, leaving source intact.
for name in list(parts):
 if 'comment' in name.lower():del parts[name];continue
 if name.endswith('.rels') or name=='[Content_Types].xml':
  tr=E.fromstring(parts[name])
  for el in list(tr):
   if any('comment' in v.lower() for v in el.attrib.values()):tr.remove(el)
  parts[name]=E.tostring(tr,xml_declaration=True,encoding='UTF-8',standalone=True)
b=io.BytesIO()
with ZipFile(b,'w',ZIP_DEFLATED) as z:
 for name,data in parts.items():z.writestr(name,data)
b.seek(0);doc=Document(b)
analysis=json.loads((R/'generated/review_analysis.json').read_text())
frozen=json.loads((R/'paper_materials/analysis/frozen_inputs.json').read_text())
labels=dict(zip(frozen['dimensions'],['Iconicity','Sensorimotor imagery','Motional salience','Facial expression','Complexity fit','Cultural familiarity','Enactment potential']))
cohorts={'de':'German','en':'English','it':'Italian','ja':'Japanese'}
def pval(v):return '< .001' if v<.001 else f'{v:.3f}'.lstrip('0')
def geometry(table,widths):
 table.autofit=False
 pr=table._tbl.tblPr
 tw=pr.find(qn('w:tblW'));tw.set(qn('w:w'),str(sum(widths)));tw.set(qn('w:type'),'dxa')
 indent=pr.find(qn('w:tblInd'))
 if indent is None:indent=OxmlElement('w:tblInd');pr.append(indent)
 indent.set(qn('w:w'),'100');indent.set(qn('w:type'),'dxa')
 grid=table._tbl.tblGrid
 for ch in list(grid):grid.remove(ch)
 for wd in widths:
  col=OxmlElement('w:gridCol');col.set(qn('w:w'),str(wd));grid.append(col)
 for row in table.rows:
  for c,wd in zip(row.cells,widths):
   c.width=Inches(wd/1440);tcw=c._tc.get_or_add_tcPr().find(qn('w:tcW'));tcw.set(qn('w:w'),str(wd));tcw.set(qn('w:type'),'dxa')
   mar=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),'75' if side in ['top','bottom'] else '100');e.set(qn('w:type'),'dxa');mar.append(e)
   c._tc.get_or_add_tcPr().append(mar)
  for p in row.cells[0].paragraphs:p.paragraph_format.keep_with_next=False
 header=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(header)
 for row in table.rows:
  for cell in row.cells:
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.0
    for run in p.runs:run.font.size=Pt(9)
def fill_table(table,headers,rows,widths):
 while len(table.rows)>1:table._tbl.remove(table.rows[-1]._tr)
 while len(table.columns)<len(headers):table.add_column(Inches(.7))
 for c,v in zip(table.rows[0].cells,headers):c.text=v
 for rr in rows:
  cells=table.add_row().cells
  for c,v in zip(cells,rr):c.text=str(v)
 for c in table.rows[0].cells:
  for r in c.paragraphs[0].runs:r.bold=True
 geometry(table,widths)

# Give the longest Table 1 label enough room before its numeric mean.
geometry(doc.tables[0],[4400,850,850,1700,1050])
summary=doc.tables[7]
summary.rows[1].cells[3].text='Positive overall correspondence; ranking and absolute agreement differ'
summary.rows[4].cells[2].text='Limited descriptive support'
summary.rows[4].cells[3].text='Item-spread correlations .07–.38; midpoint-adjusted .00–.19'
table=doc.tables[2]
rows=[]
for c in analysis['contrasts']:
 rows.append([labels[c['dimension']],cohorts[c['cohort']],f"{c['b']:.2f}",f"[{c['ci_low']:.2f}, {c['ci_high']:.2f}]",pval(c['p']),pval(c['p_holm'])])
fill_table(table,['Dimension','Cohort','b','95% CI','p','Holm p'],rows,[2500,1300,650,1900,900,900])

doc.add_page_break()
doc.add_paragraph('Appendix A. Response-distribution diagnostics',style='Heading 1')
doc.add_paragraph('Table A1. Endpoint / midpoint use by dimension and survey-language cohort (%)')
doc.add_paragraph('Each cell reports the percentage of individual ratings at either endpoint (1 or 5), followed by the percentage at the midpoint (3). Each dimension has 1,460 German, 1,000 English, 1,020 Italian, and 1,120 Japanese ratings. All 4,600 retained evaluations, including controls, are included. These are descriptive distributions, not item-adjusted estimates of response style.')
lookup={(v['dimension'],v['cohort']):v for v in analysis['response_style']}
rows=[]
for dim,label in labels.items():
 vals=[]
 for co in cohorts:
  v=lookup[(dim,co)];vals.append(f"{v['extreme']*100:.1f} / {v['midpoint']*100:.1f}")
 rows.append([label]+vals)
ta=doc.add_table(rows=1,cols=5);ta.style=doc.tables[0].style
fill_table(ta,['Dimension','German','English','Italian','Japanese'],rows,[2850,1325,1325,1325,1325])
doc.add_paragraph('The participant workbook used for these calculations was verified against the SHA-256 recorded in the existing analysis audit; it contains 4,600 retained evaluations from 230 participants. No participant identifiers or trial-level rows are reproduced here.')

doc.add_paragraph('Appendix B. Model provenance and parsing audit',style='Heading 1')
doc.add_paragraph('Table B1. Saved model identifiers, access routes, and evaluation dates')
doc.add_paragraph('“Not retained” means the exact evaluation date cannot be recovered from the retained analysis package; it is not inferred from a directory name or file modification time. Identifiers are copied verbatim from the frozen model metadata. A provider alias does not establish an immutable model snapshot. The native Gemini routes and auxiliary frame route are documented by the saved inference scripts; exact requested identifiers remain unavailable for these three configurations.')
rows=[]
for key,m in frozen['models'].items():
 ids=', '.join(m['logged_model_ids']) or 'Not logged'
 route='Google native API' if 'native-2fps' in key else ('Google native API (frames)' if key=='gemini-3.8-flash-2fps' else 'OpenRouter')
 date='19 Oct 2026' if key.startswith('elorian') else 'Not retained'
 rows.append([m['label']+'; '+m['input'],ids,route,date])
tb=doc.add_table(rows=1,cols=4);tb.style=doc.tables[0].style
fill_table(tb,['Configuration','Saved exact ID','Access route','Date'],rows,[3000,2700,1450,1000])
doc.add_paragraph('Glimpse saved canonical version: elorian/glimpse-20261016. Model family, API route, sampling, and image preprocessing vary together; these configurations cannot support a controlled architecture-only comparison.')
doc.add_paragraph('Table B2. Final completeness and documented field remapping')
rows=[]
for key,m in frozen['models'].items():
 re=analysis['remaps'][key]
 scores=frozen['scores'][key]
 complete=sum(len(v)==7 and all(isinstance(s,(int,float)) and 1<=s<=5 for s in v) for v in scores.values())
 assert complete==149,(key,complete)
 rows.append([m['label']+'; '+m['input'],f'0 / {complete}',f"{re['items']} / 149 ({100*re['items']/149:.1f}%)",str(re['fields'])])
tc=doc.add_table(rows=1,cols=4);tc.style=doc.tables[0].style
fill_table(tc,['Configuration','Missing final vectors','Pairings remapped','Fields'],rows,[3900,1300,1900,1050])
doc.add_paragraph('Final completeness is checked on the frozen seven-score vectors for all 149 model pairings (134 standard and 15 controls); 14 controls have a human reference. Remap counts come from field_normalization.json. Zero documented remaps does not imply zero initial parsing failures. Historical parse-failure and retry rates cannot be established uniformly from the retained records and are therefore not reported as zero.')
doc.core_properties.comments=''
doc.core_properties.last_modified_by=''
doc.save(OUT)
with ZipFile(OUT) as z:
 for name in z.namelist():
  if name.startswith('word/') and name.endswith('.xml'):
   tree=E.fromstring(z.read(name))
   assert not tree.xpath('//w:ins | //w:del | //w:commentRangeStart | //w:commentReference | //w:trackRevisions',namespaces=NS),name
 assert not any('comment' in n.lower() for n in z.namelist())
print('FINAL',OUT)
